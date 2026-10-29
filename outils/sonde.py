"""Reconnaissance des sources de prix (lancée à la main, une fois) : ce que renvoie chaque source, pour écrire le
relevé. N'écrit rien dans le dépôt."""
import io
import re
import urllib.request

UA = {"User-Agent": "Mozilla/5.0 (Valeur-engrais; releve hebdomadaire des cours)"}


def lire(url):
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40) as r:
            d = r.read()
            print(f"--- {url} : {r.status}, {len(d)} octets, {r.headers.get('Content-Type')}")
            return d
    except Exception as e:
        print(f"--- {url} : ÉCHEC {e}")
        return b""


# 1. robots.txt des sites de cotations
for site in ("https://www.web-agri.fr", "https://www.terre-net.fr", "https://www.agri-mutuel.com", "https://www.pleinchamp.com"):
    t = lire(site + "/robots.txt").decode("utf-8", "replace")
    print("\n".join(l for l in t.splitlines() if l.lower().startswith(("user-agent", "disallow", "allow")))[:1500])

# 2. les pages de cotations : ce qui entoure les noms d'engrais
for url in ("https://www.web-agri.fr/marches-agricoles/engrais", "https://www.terre-net.fr/marche-agricole/engrais",
            "https://www.agri-mutuel.com/cotations/engrais/", "https://www.pleinchamp.com/les-cours/intrants"):
    h = lire(url).decode("utf-8", "replace")
    txt = re.sub(r"<script.*?</script>|<style.*?</style>", " ", h, flags=re.S)
    txt = re.sub(r"<[^>]+>", "|", txt)
    txt = re.sub(r"\s+", " ", txt)
    txt = re.sub(r"(\| ?)+", "|", txt)
    vus = 0
    for m in re.finditer(r"(Ammonitrate|Solution azot|Urée|Super ?45|TSP|Chlorure|Potass|Phosphat|DAP|18-46)", txt, re.I):
        print("   …", txt[max(0, m.start() - 60): m.start() + 200])
        vus += 1
        if vus > 25:
            break
    print("   (fin de", url, ")")

# 3. Banque mondiale : le lien du classeur mensuel
h = lire("https://www.worldbank.org/en/research/commodity-markets").decode("utf-8", "replace")
liens = sorted(set(re.findall(r'https?://[^"\'\s]+?CMO-Historical-Data-Monthly\.xlsx', h)))
print("liens classeur :", liens)
if liens:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(lire(liens[0])), read_only=True, data_only=True)
    print("feuilles :", wb.sheetnames)
    lignes = list(wb["Monthly Prices"].iter_rows(values_only=True))
    for i, row in enumerate(lignes[:8]):
        print(i, list(row)[:80])
    print("dernières :", [list(r)[:1] for r in lignes[-3:]])

# 4. INSEE : flux IPAMPA et la série des engrais azotés simples
fl = lire("https://bdm.insee.fr/series/sdmx/dataflow").decode("utf-8", "replace")
flux = sorted(set(re.findall(r'id="([^"]*IPAMPA[^"]*)"', fl)))
print("flux :", flux)
print(lire("https://bdm.insee.fr/series/sdmx/data/SERIES_BDM/010777321?lastNObservations=3").decode("utf-8", "replace")[:2500])
for f in flux:
    d = lire(f"https://bdm.insee.fr/series/sdmx/data/{f}?lastNObservations=1").decode("utf-8", "replace")
    for t in re.finditer(r"<Series [^>]*>", d):
        if re.search(r"engrais|potass|phosph|azot", t.group(0), re.I):
            print("  ", t.group(0)[:700])
